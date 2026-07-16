"""Smoke tests for mteb_eval toolkit."""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest

from mteb_eval.cache import configure_cache
from mteb_eval.languages import ML16_LANGUAGES, resolve_languages
from mteb_eval.runner import release_task_memory
from mteb_eval.model_loader import (
    build_hf_model_kwargs,
    configure_max_seq_len,
    resolve_model_source,
    resolve_torch_dtype,
    validate_local_checkpoint,
)
from mteb_eval.prompts import configure_prompt_prefixes
from mteb_eval.summary import (
    build_language_summary_rows,
    build_summary_rows,
    merge_shard_results,
    print_summary,
    write_language_detail_csv,
    write_language_outputs,
    write_language_summary_csv,
    write_summary_csv,
)
from mteb_eval.tasks import (
    CLF_CLUST_RERANK_TYPES,
    ML16_CLF_CLUST_RERANK_MANIFEST,
    RETRIEVAL_FAST_OMITTED,
    RETRIEVAL_FAST_TASKS,
    WEBLINX_RERANKING,
    expand_task_types,
    expected_task_names,
    filter_tasks_by_languages,
    is_clf_clust_rerank_types,
    list_mteb_eval_all_task_names,
    load_manifest,
    partition_task_names,
    resolve_task_names,
    resolve_task_selection,
    resolve_tasks,
    validate_against_manifest,
)
from mteb.results.model_result import ModelResult
from mteb.results.task_result import TaskResult


def _sample_task_result(name: str = "BIOSSES") -> TaskResult:
    return TaskResult(
        task_name=name,
        dataset_revision="rev",
        mteb_version="2.18.0",
        evaluation_time=1.0,
        scores={
            "test": [
                {
                    "main_score": 0.1,
                    "hf_subset": "default",
                    "languages": ["eng-Latn"],
                }
            ]
        },
    )


def test_manifest_has_19_tasks():
    manifest = load_manifest()
    assert manifest["expected_count"] == 19
    assert len(manifest["tasks"]) == 19


def test_resolve_sts_retrieval_task_count():
    tasks = resolve_tasks()
    assert len(tasks) == 19
    validate_against_manifest(tasks)


def test_expected_task_names_match_plan():
    names = set(expected_task_names())
    assert "FEVERHardNegatives" in names
    assert "HotpotQAHardNegatives" in names
    assert "BIOSSES" in names
    assert "STS22.v2" in names
    assert len(names) == 19


def test_cache_dir_sets_hf_env():
    with tempfile.TemporaryDirectory() as tmp:
        for key in (
            "HF_HOME",
            "HF_HUB_CACHE",
            "HF_DATASETS_CACHE",
            "TRANSFORMERS_CACHE",
            "HF_DATASETS_OFFLINE",
            "HF_HUB_OFFLINE",
            "TRANSFORMERS_OFFLINE",
        ):
            os.environ.pop(key, None)

        cfg = configure_cache(cache_dir=tmp, offline=True)
        assert cfg.cache_dir == Path(tmp).resolve()
        assert os.environ["HF_HOME"] == str(Path(tmp).resolve())
        assert os.environ["HF_HUB_CACHE"].endswith("/hub")
        assert os.environ["HF_DATASETS_OFFLINE"] == "1"
        assert os.environ["HF_HUB_OFFLINE"] == "1"


def test_default_cache_does_not_set_hf_home():
    saved = os.environ.get("HF_HOME")
    os.environ.pop("HF_HOME", None)
    try:
        cfg = configure_cache(default_cache=True)
        assert cfg.use_default_cache is True
        assert "HF_HOME" not in os.environ
    finally:
        if saved is not None:
            os.environ["HF_HOME"] = saved


def test_resolve_local_model_path(tmp_path: Path):
    model_dir = tmp_path / "mock-st"
    model_dir.mkdir()
    (model_dir / "config.json").write_text(json.dumps({"model_type": "bert"}))
    (model_dir / "modules.json").write_text("{}")

    source = resolve_model_source(model_path=str(model_dir), model=None, hub_id=None)
    assert source.is_local is True
    assert source.path == str(model_dir.resolve())


def test_resolve_model_path_convenience_dir(tmp_path: Path):
    model_dir = tmp_path / "checkpoint"
    model_dir.mkdir()
    (model_dir / "config_sentence_transformers.json").write_text("{}")

    source = resolve_model_source(model=str(model_dir), model_path=None, hub_id=None)
    assert source.is_local is True


def test_model_and_model_path_mutually_exclusive(tmp_path: Path):
    model_dir = tmp_path / "m"
    model_dir.mkdir()
    (model_dir / "modules.json").write_text("{}")
    with pytest.raises(ValueError, match="not both"):
        resolve_model_source(model="org/model", model_path=str(model_dir), hub_id=None)


def test_missing_model_source_raises():
    with pytest.raises(ValueError, match="Provide --model"):
        resolve_model_source(model=None, model_path=None, hub_id=None)


def test_validate_local_checkpoint_st_layout(tmp_path: Path):
    model_dir = tmp_path / "st"
    model_dir.mkdir()
    (model_dir / "modules.json").write_text("{}")
    layout = validate_local_checkpoint(model_dir)
    assert layout.kind == "sentence-transformers"


def test_validate_local_checkpoint_transformers_layout(tmp_path: Path):
    model_dir = tmp_path / "tf"
    model_dir.mkdir()
    (model_dir / "config.json").write_text("{}")
    (model_dir / "model.safetensors").write_bytes(b"")
    layout = validate_local_checkpoint(model_dir)
    assert layout.kind == "transformers"


def test_print_summary_includes_failed_tasks(capsys):
    task_result = MagicMock()
    task_result.task_name = "BIOSSES"
    task_result.get_score.return_value = 0.5
    results = SimpleNamespace(task_results=[task_result])
    print_summary(
        results,
        {"BIOSSES": 1.0, "STS12": 0.5},
        failures={"STS12": "dataset not found"},
    )
    out = capsys.readouterr().out
    assert "BIOSSES" in out
    assert "0.5000" in out
    assert "STS12" in out
    assert "FAILED" in out
    assert "dataset not found" in out
    assert "AVERAGE" in out
    assert "Failed tasks: 1" in out


def test_continue_on_error_collects_failures(tmp_path: Path):
    from mteb_eval.runner import EvalRunResult

    ok_task = MagicMock()
    ok_task.metadata.name = "BIOSSES"
    bad_task = MagicMock()
    bad_task.metadata.name = "STS12"

    ok_result = EvalRunResult(
        model_result=ModelResult(
            model_name="mteb/baseline-random-encoder",
            model_revision=None,
            task_results=[_sample_task_result("BIOSSES")],
            exceptions=None,
        ),
        timings={"BIOSSES": 1.0, "STS12": 0.5},
        failures={"STS12": "boom"},
        task_prompts={},
        exit_code=1,
        model_name="mteb/baseline-random-encoder",
    )

    args = SimpleNamespace(
        cache_dir=None,
        default_cache=True,
        model="mteb/baseline-random-encoder",
        model_path=None,
        hub_id=None,
        model_type="auto",
        offline=False,
        benchmark="MTEB(eng, v2)",
        task_types=["STS"],
        tasks=None,
        tasks_preset=None,
        languages=None,
        languages_preset="ml16",
        exclusive_language_filter=False,
        output_dir=str(tmp_path / "out"),
        device="cpu",
        dtype="bfloat16",
        attn_implementation="sdpa",
        batch_size=8,
        query_batch_size=None,
        corpus_batch_size=None,
        overwrite="only-missing",
        continue_on_error=True,
        max_seq_len=512,
        query_prefix=None,
        document_prefix=None,
        verbose=False,
    )

    with patch("mteb_eval.evaluate.run_evaluation", return_value=ok_result) as mock_run, patch(
        "mteb_eval.evaluate.build_parser"
    ) as mock_parser:
        from mteb_eval import evaluate as evaluate_module

        mock_parser.return_value.parse_args.return_value = args
        exit_code = evaluate_module.main([])

    assert exit_code == 1
    mock_run.assert_called_once()
    assert (tmp_path / "out" / "summary.csv").exists()


def test_write_summary_csv_includes_average(tmp_path: Path):
    rows = [
        {
            "task": "STS12",
            "score": 0.8,
            "time_s": 2.0,
            "status": "ok",
            "query_prompt": "q1",
            "document_prompt": "d1",
            "error": "",
        },
        {
            "task": "STS13",
            "score": 0.6,
            "time_s": 4.0,
            "status": "ok",
            "query_prompt": "q2",
            "document_prompt": "d2",
            "error": "",
        },
    ]
    csv_path = tmp_path / "summary.csv"
    write_summary_csv(csv_path, rows)
    content = csv_path.read_text(encoding="utf-8")
    assert "AVERAGE" in content
    assert "0.700000" in content
    assert "3.000" in content


def test_configure_prompt_prefixes_on_sentence_transformer():
    st_model = SimpleNamespace(prompts={"query": "old-q", "document": "old-d"})
    wrapper = SimpleNamespace(model=st_model, model_prompts={"query": "old-q"})

    configure_prompt_prefixes(
        wrapper,
        query_prefix="query: ",
        document_prefix="document: ",
    )
    assert st_model.prompts["query"] == "query: "
    assert st_model.prompts["document"] == "document: "
    assert wrapper.model_prompts["query"] == "query: "
    assert wrapper.model_prompts["document"] == "document: "


def test_configure_max_seq_len_sentence_transformer():
    st_model = SimpleNamespace(max_seq_length=8192)
    wrapper = SimpleNamespace(model=st_model)

    configure_max_seq_len(wrapper, 512)
    assert st_model.max_seq_length == 512


def test_configure_max_seq_len_eurobert_wrapper():
    from mteb_eval.model_loader import EuroBertEncoderWrapper

    model = EuroBertEncoderWrapper(model=MagicMock(), tokenizer=MagicMock(), device="cpu")
    model.max_length = 2048
    configure_max_seq_len(model, 512)
    assert model.max_length == 512


def test_release_task_memory_runs():
    release_task_memory()


def test_ml16_preset_has_16_codes():
    assert len(ML16_LANGUAGES) == 16
    assert "eng-Latn" in ML16_LANGUAGES
    assert "jpn-Jpan" in ML16_LANGUAGES
    assert "zho-Hans" in ML16_LANGUAGES


def test_resolve_languages_preset():
    langs = resolve_languages(languages_preset="ml16")
    assert langs == list(ML16_LANGUAGES)


def test_resolve_languages_explicit_overrides_preset():
    langs = resolve_languages(languages=["eng-Latn"], languages_preset="ml16")
    assert langs == ["eng-Latn"]


def test_resolve_languages_none_disables_filter():
    assert resolve_languages(languages_preset="none") is None
    assert resolve_languages(languages_preset=None) is None


def test_resolve_languages_default_is_ml16():
    langs = resolve_languages()
    assert langs == list(ML16_LANGUAGES)


def test_filter_tasks_by_languages_sts17():
    import mteb

    task = mteb.get_task("STS17")
    before = len(task.hf_subsets)
    filtered = filter_tasks_by_languages([task], list(ML16_LANGUAGES))
    assert len(filtered) == 1
    assert len(filtered[0].hf_subsets) == before


def test_filter_skips_no_overlap_task():
    import mteb

    task = mteb.get_task("TwitterHjerneRetrieval")
    filtered = filter_tasks_by_languages([task], list(ML16_LANGUAGES))
    assert filtered == []


def test_partition_task_names_round_robin():
    names = ["C", "A", "B", "D", "E"]
    parts = partition_task_names(names, 3)
    assert parts == [["A", "D"], ["B", "E"], ["C"]]


def test_merge_shard_results(tmp_path: Path):
    shard_a = tmp_path / "shard0"
    shard_b = tmp_path / "shard1"
    shard_a.mkdir()
    shard_b.mkdir()

    result_a = _sample_task_result("STS12")
    result_b = _sample_task_result("STS13")
    model_dump_a = {
        "model_name": "test/model",
        "model_revision": None,
        "task_results": [result_a.model_dump()],
        "exceptions": None,
    }
    model_dump_b = {
        "model_name": "test/model",
        "model_revision": None,
        "task_results": [result_b.model_dump()],
        "exceptions": None,
    }
    (shard_a / "summary.json").write_text(json.dumps(model_dump_a), encoding="utf-8")
    (shard_b / "summary.json").write_text(json.dumps(model_dump_b), encoding="utf-8")
    (shard_a / "run_meta.json").write_text(
        json.dumps({"timings": {"STS12": 1.0}, "failures": {}, "task_prompts": {}, "exit_code": 0}),
        encoding="utf-8",
    )
    (shard_b / "run_meta.json").write_text(
        json.dumps({"timings": {"STS13": 2.0}, "failures": {}, "task_prompts": {}, "exit_code": 0}),
        encoding="utf-8",
    )

    output_dir = tmp_path / "merged"
    merged = merge_shard_results([shard_a, shard_b], output_dir)
    assert len(merged.model_result.task_results) == 2
    assert merged.timings["STS12"] == 1.0
    assert merged.timings["STS13"] == 2.0
    assert (output_dir / "summary.json").exists()
    assert (output_dir / "summary.csv").exists() is False  # CSV written by caller


def test_retrieval_fast_preset_membership():
    assert len(RETRIEVAL_FAST_TASKS) == 12
    assert len(RETRIEVAL_FAST_OMITTED) == 6
    assert set(RETRIEVAL_FAST_TASKS).isdisjoint(RETRIEVAL_FAST_OMITTED)
    assert "ArguAna" in RETRIEVAL_FAST_TASKS
    assert "BelebeleRetrieval" in RETRIEVAL_FAST_OMITTED
    assert "MIRACLRetrievalHardNegatives" in RETRIEVAL_FAST_OMITTED


def test_resolve_task_names_preset():
    names, from_preset = resolve_task_names(tasks_preset="retrieval-fast")
    assert from_preset is True
    assert names == list(RETRIEVAL_FAST_TASKS)


def test_resolve_task_names_mutually_exclusive():
    with pytest.raises(ValueError, match="not both"):
        resolve_task_names(tasks=["ArguAna"], tasks_preset="retrieval-fast")


def test_resolve_tasks_retrieval_fast_on_multilingual():
    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=["Retrieval"],
        task_names=list(RETRIEVAL_FAST_TASKS),
        allow_missing_task_names=True,
    )
    names = {t.metadata.name for t in tasks}
    assert names == set(RETRIEVAL_FAST_TASKS)
    assert names.isdisjoint(RETRIEVAL_FAST_OMITTED)


def test_build_hf_model_kwargs_auto_empty():
    assert build_hf_model_kwargs(dtype="auto") == {}


def test_build_hf_model_kwargs_bfloat16():
    import torch

    kwargs = build_hf_model_kwargs(dtype="bfloat16")
    assert kwargs == {"torch_dtype": torch.bfloat16}


def test_build_hf_model_kwargs_attn_only_when_set():
    import torch

    assert "attn_implementation" not in build_hf_model_kwargs(dtype="auto")
    kwargs = build_hf_model_kwargs(dtype="float16", attn_implementation="sdpa")
    assert kwargs["torch_dtype"] == torch.float16
    assert kwargs["attn_implementation"] == "sdpa"


def test_resolve_torch_dtype_auto_is_none():
    assert resolve_torch_dtype("auto") is None


def test_ensure_mteb_model_meta_attaches_qwen3_registry_entry():
    from mteb_eval.model_loader import ensure_mteb_model_meta

    model = SimpleNamespace(mteb_model_meta=None)
    ensure_mteb_model_meta(model, hub_id="Qwen/Qwen3-Embedding-0.6B")
    assert model.mteb_model_meta is not None
    assert model.mteb_model_meta.name == "Qwen/Qwen3-Embedding-0.6B"
    assert model.mteb_model_meta.revision
    assert model.mteb_model_meta.revision != "no_revision_available"


def test_ensure_mteb_model_meta_skips_when_complete():
    from mteb_eval.model_loader import ensure_mteb_model_meta

    existing = SimpleNamespace(name="already/set", revision="abc123")
    model = SimpleNamespace(mteb_model_meta=existing)
    ensure_mteb_model_meta(model, hub_id="Qwen/Qwen3-Embedding-0.6B")
    assert model.mteb_model_meta is existing


def test_load_qwen3_local_attaches_model_meta():
    """Regression: missing mteb_model_meta made every task fail with Path / None."""
    from mteb_eval.model_loader import ModelSource, load_qwen3_local

    source = ModelSource(
        path="/tmp/fake-qwen3",
        is_local=True,
        hub_id="Qwen/Qwen3-Embedding-0.6B",
    )
    fake_model = SimpleNamespace(mteb_model_meta=None)

    with patch(
        "mteb.models.model_implementations.qwen3_models.q3e_instruct_loader",
        return_value=fake_model,
    ) as mock_loader:
        model = load_qwen3_local(
            source,
            device="cpu",
            hub_id="Qwen/Qwen3-Embedding-0.6B",
            model_kwargs={},
        )

    assert model is fake_model
    assert model.mteb_model_meta is not None
    assert model.mteb_model_meta.name == "Qwen/Qwen3-Embedding-0.6B"
    assert model.mteb_model_meta.revision
    args, kwargs = mock_loader.call_args
    assert args[0] == "/tmp/fake-qwen3"
    revision_arg = kwargs.get("revision")
    assert revision_arg is not None
    assert isinstance(revision_arg, str)


def test_load_embedding_model_passes_nested_model_kwargs():
    from mteb_eval.model_loader import ModelSource, load_embedding_model

    source = ModelSource(path="org/demo-model", is_local=False, hub_id="org/demo-model")
    with (
        patch("mteb.get_model_meta", return_value=None),
        patch("sentence_transformers.SentenceTransformer") as mock_st,
    ):
        mock_st.return_value = MagicMock()
        load_embedding_model(source, dtype="bfloat16", attn_implementation="sdpa")

    _, kwargs = mock_st.call_args
    import torch

    assert kwargs["model_kwargs"]["torch_dtype"] == torch.bfloat16
    assert kwargs["model_kwargs"]["attn_implementation"] == "sdpa"
    assert "torch_dtype" not in kwargs


def test_language_summary_rows_and_csv(tmp_path: Path):
    from mteb.results.model_result import ModelResult
    from mteb.results.task_result import TaskResult

    tr = TaskResult(
        task_name="STS17",
        dataset_revision="rev",
        mteb_version="2.18.0",
        evaluation_time=1.0,
        scores={
            "test": [
                {"main_score": 0.8, "hf_subset": "en-en", "languages": ["eng-Latn"]},
                {"main_score": 0.6, "hf_subset": "en-de", "languages": ["eng-Latn", "deu-Latn"]},
                {"main_score": 0.4, "hf_subset": "de-de", "languages": ["deu-Latn"]},
            ]
        },
    )
    results = ModelResult(
        model_name="test/model",
        model_revision=None,
        task_results=[tr],
        exceptions=None,
    )
    rows = build_language_summary_rows(results)
    by_lang = {r["language"]: r for r in rows}
    assert by_lang["eng-Latn"]["mean_score"] == pytest.approx(0.7)
    assert by_lang["deu-Latn"]["mean_score"] == pytest.approx(0.5)
    assert by_lang["eng-Latn"]["n_scores"] == 2

    out = write_language_outputs(tmp_path, results)
    assert out is not None
    lang_csv, detail_csv = out
    assert lang_csv.exists()
    assert detail_csv.exists()
    lang_text = lang_csv.read_text(encoding="utf-8")
    assert "eng-Latn" in lang_text
    assert "AVERAGE" in lang_text
    detail_text = detail_csv.read_text(encoding="utf-8")
    assert "en-de" in detail_text
    assert detail_text.count("eng-Latn") >= 2


def test_resolve_qrels_config_prefers_qrels():
    from mteb_eval.offline_compat import resolve_qrels_config

    assert (
        resolve_qrels_config(["corpus", "qrels", "queries"], subset_config=None)
        == "qrels"
    )
    assert (
        resolve_qrels_config(["default", "corpus", "queries"], subset_config=None)
        == "default"
    )
    assert (
        resolve_qrels_config(
            ["default", "corpus", "qrels", "queries"], subset_config=None
        )
        == "qrels"
    )
    assert resolve_qrels_config(["x"], subset_config="en") == "en-qrels"


def test_offline_compat_loads_qrels_when_default_advertised(tmp_path: Path):
    """Reproduce: hub lists 'default' but only corpus/qrels/queries are cached."""
    import os
    from unittest.mock import patch

    from mteb_eval.cache import configure_cache
    from mteb_eval.offline_compat import apply_mteb_offline_compat

    configure_cache(cache_dir=tmp_path / "hf", offline=False)
    apply_mteb_offline_compat()

    import mteb
    from mteb.abstasks.retrieval_dataset_loaders import RetrievalDatasetLoader

    # Warm arrow cache online for a qrels-layout dataset
    task = mteb.get_task("WinoGrande")
    path = task.metadata.dataset["path"]
    rev = task.metadata.dataset["revision"]
    task.load_data()

    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["HF_DATASETS_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"

    with patch(
        "mteb.abstasks.retrieval_dataset_loaders.get_dataset_config_names",
        return_value=["default", "corpus", "qrels", "queries"],
    ):
        data = RetrievalDatasetLoader(
            hf_repo=path, revision=rev, split="test", config="default"
        ).load()
    assert data["relevant_docs"]
    assert len(data["corpus"]) > 0
    assert len(data["queries"]) > 0


def test_offline_compat_still_loads_arguana_default(tmp_path: Path):
    import os

    from mteb_eval.cache import configure_cache
    from mteb_eval.offline_compat import apply_mteb_offline_compat

    configure_cache(cache_dir=tmp_path / "hf", offline=False)
    apply_mteb_offline_compat()

    import mteb

    task = mteb.get_task("ArguAna")
    task.load_data()

    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["HF_DATASETS_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"

    task2 = mteb.get_task("ArguAna")
    task2.load_data()
    assert task2.data_loaded


def test_evaluate_cli_no_continue_on_error():
    from mteb_eval.evaluate import build_parser

    args = build_parser().parse_args(
        [
            "--default-cache",
            "--model",
            "org/model",
            "--output-dir",
            "/tmp/out",
            "--no-continue-on-error",
        ]
    )
    assert args.continue_on_error is False


def test_prefetch_cli_defaults():
    from mteb_eval.prefetch import build_parser
    from mteb_eval.languages import languages_from_args

    args = build_parser().parse_args(["--default-cache"])
    assert args.benchmark == "MTEB(Multilingual, v2)"
    assert args.languages_preset == "ml16"
    langs = languages_from_args(args)
    assert langs is not None and len(langs) == 16


def test_ml16_filters_belebele_subsets():
    from mteb_eval.languages import ML16_LANGUAGES
    from mteb_eval.tasks import resolve_tasks

    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=["Retrieval"],
        task_names=["BelebeleRetrieval"],
        languages=list(ML16_LANGUAGES),
    )
    assert len(tasks) == 1
    assert len(tasks[0].hf_subsets) < 376
    assert len(tasks[0].hf_subsets) == 261


def test_default_cli_languages_filter_belebele():
    """Reproduce: omitting --languages-preset still applies ml16."""
    from mteb_eval.evaluate import build_parser
    from mteb_eval.languages import languages_from_args
    from mteb_eval.tasks import resolve_tasks

    args = build_parser().parse_args(
        [
            "--default-cache",
            "--model",
            "org/model",
            "--output-dir",
            "/tmp/out",
            "--tasks",
            "BelebeleRetrieval",
            "--task-types",
            "Retrieval",
        ]
    )
    languages = languages_from_args(args)
    assert languages is not None and len(languages) == 16
    tasks = resolve_tasks(
        benchmark=args.benchmark,
        task_types=args.task_types,
        task_names=args.tasks,
        languages=languages,
        exclusive_language_filter=args.exclusive_language_filter,
    )
    assert len(tasks) == 1
    assert len(tasks[0].hf_subsets) == 261


def test_languages_preset_none_keeps_all_belebele():
    from mteb_eval.evaluate import build_parser
    from mteb_eval.languages import languages_from_args
    from mteb_eval.tasks import resolve_tasks

    args = build_parser().parse_args(
        [
            "--default-cache",
            "--model",
            "org/model",
            "--output-dir",
            "/tmp/out",
            "--tasks",
            "BelebeleRetrieval",
            "--task-types",
            "Retrieval",
            "--languages-preset",
            "none",
        ]
    )
    languages = languages_from_args(args)
    assert languages is None
    tasks = resolve_tasks(
        benchmark=args.benchmark,
        task_types=args.task_types,
        task_names=args.tasks,
        languages=languages,
    )
    assert len(tasks) == 1
    assert len(tasks[0].hf_subsets) == 376


def test_expand_task_types_includes_hierarchical_clustering():
    assert expand_task_types(["Clustering"]) == [
        "Clustering",
        "HierarchicalClustering",
    ]
    assert expand_task_types(["Classification", "Clustering", "Reranking"]) == [
        "Classification",
        "Clustering",
        "HierarchicalClustering",
        "Reranking",
    ]
    assert expand_task_types(["STS", "Retrieval"]) == ["STS", "Retrieval"]


def test_is_clf_clust_rerank_types():
    assert is_clf_clust_rerank_types(CLF_CLUST_RERANK_TYPES)
    assert is_clf_clust_rerank_types(["Reranking", "Classification", "Clustering"])
    assert not is_clf_clust_rerank_types(["STS", "Retrieval"])
    assert not is_clf_clust_rerank_types(["Classification", "Clustering"])


def test_ml16_clf_clust_rerank_manifest():
    manifest = load_manifest(ML16_CLF_CLUST_RERANK_MANIFEST)
    assert manifest["benchmark"] == "MTEB(Multilingual, v2)"
    assert manifest["languages_preset"] == "ml16"
    assert manifest["expected_count"] == 38
    assert manifest["counts_by_type"] == {
        "Classification": 21,
        "Clustering": 12,
        "Reranking": 5,
    }
    assert len(manifest["tasks"]) == 38


def test_resolve_ml16_clf_clust_rerank_matches_manifest():
    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=list(CLF_CLUST_RERANK_TYPES),
        languages=list(ML16_LANGUAGES),
    )
    assert len(tasks) == 38
    validate_against_manifest(tasks, manifest_name=ML16_CLF_CLUST_RERANK_MANIFEST)
    by_type: dict[str, int] = {}
    for task in tasks:
        by_type[task.metadata.type] = by_type.get(task.metadata.type, 0) + 1
    assert by_type == {"Classification": 21, "Clustering": 12, "Reranking": 5}


def test_ml16_shrinks_massive_intent_and_sib200():
    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=["Classification", "Clustering"],
        task_names=["MassiveIntentClassification", "SIB200ClusteringS2S"],
        languages=list(ML16_LANGUAGES),
    )
    by_name = {t.metadata.name: t for t in tasks}
    assert len(by_name["MassiveIntentClassification"].hf_subsets) == 15
    assert len(by_name["SIB200ClusteringS2S"].hf_subsets) == 15


def test_ml16_skips_non_overlapping_classification():
    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=["Classification"],
        task_names=["DalajClassification", "FinancialPhrasebankClassification"],
        languages=list(ML16_LANGUAGES),
    )
    names = {t.metadata.name for t in tasks}
    assert names == {"FinancialPhrasebankClassification"}


def test_evaluate_cli_defaults_remain_sts_retrieval():
    from mteb_eval.evaluate import build_parser

    args = build_parser().parse_args(
        ["--default-cache", "--model", "org/model", "--output-dir", "/tmp/out"]
    )
    assert args.task_types == ["STS", "Retrieval"]


def test_prefetch_auto_validates_clf_clust_rerank_ml16():
    from mteb_eval.prefetch import _resolve_manifest_validation

    manifest = _resolve_manifest_validation(
        should_validate=None,
        benchmark="MTEB(Multilingual, v2)",
        task_types=["Classification", "Clustering", "Reranking"],
        names=None,
        languages=list(ML16_LANGUAGES),
        languages_preset="ml16",
    )
    assert manifest == ML16_CLF_CLUST_RERANK_MANIFEST

    # Default STS+Retrieval should not auto-validate the clf manifest.
    assert (
        _resolve_manifest_validation(
            should_validate=None,
            benchmark="MTEB(Multilingual, v2)",
            task_types=["STS", "Retrieval"],
            names=None,
            languages=list(ML16_LANGUAGES),
            languages_preset="ml16",
        )
        is None
    )


def test_mteb_eval_all_preset_membership():
    names = list_mteb_eval_all_task_names("MTEB(Multilingual, v2)")
    assert WEBLINX_RERANKING not in names
    assert "ArguAna" in names
    assert "BelebeleRetrieval" not in names
    assert any("STS" in n or n.startswith("STS") or n.endswith("STS") or "STS" == n for n in names) or any(
        n in names for n in ("STS12", "STS13", "STS14", "STS15", "STS17", "STS22.v2", "STSBenchmark", "SICK-R")
    )
    assert "MassiveIntentClassification" in names
    assert "StackExchangeClustering.v2" in names
    assert "AlloprofReranking" in names

    types, sel_names, from_preset = resolve_task_selection(
        benchmark="MTEB(Multilingual, v2)",
        task_types=["STS", "Retrieval"],
        tasks_preset="mteb-eval-all",
    )
    assert from_preset is True
    assert set(types) == {
        "STS",
        "Classification",
        "Clustering",
        "Reranking",
        "Retrieval",
    }
    assert sel_names is not None
    assert set(sel_names) == set(names)


def test_mteb_eval_all_ml16_counts():
    types, names, _ = resolve_task_selection(
        benchmark="MTEB(Multilingual, v2)",
        task_types=None,
        tasks_preset="mteb-eval-all",
    )
    tasks = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=types,
        task_names=names,
        languages=list(ML16_LANGUAGES),
        allow_missing_task_names=True,
    )
    by_type: dict[str, int] = {}
    for task in tasks:
        by_type[task.metadata.type] = by_type.get(task.metadata.type, 0) + 1
    assert WEBLINX_RERANKING not in {t.metadata.name for t in tasks}
    assert by_type.get("Classification") == 21
    assert by_type.get("Clustering") == 12
    assert by_type.get("Reranking") == 4
    assert by_type.get("Retrieval") == 12
    assert by_type.get("STS") == 13
    assert len(tasks) == 62
    assert {t.metadata.name for t in tasks if t.metadata.type == "Retrieval"} <= set(
        RETRIEVAL_FAST_TASKS
    )


def test_parallel_worker_keeps_preset_task_types():
    """Regression: evaluate_parallel workers must use coordinator-resolved types.

    With mteb-eval-all, CLI defaults remain STS/Retrieval. If workers keep those
    defaults, Classification/Clustering names raise before summary.json is written
    ("Missing shard summary").
    """
    from mteb_eval.tasks import DEFAULT_TASK_TYPES, partition_task_names

    types, names, from_preset = resolve_task_selection(
        benchmark="MTEB(Multilingual, v2)",
        task_types=list(DEFAULT_TASK_TYPES),
        tasks_preset="mteb-eval-all",
    )
    assert from_preset is True
    assert set(types) >= {"Classification", "Clustering", "Reranking", "STS", "Retrieval"}
    assert names is not None

    partition = partition_task_names(names, 8)[0]
    assert partition

    with pytest.raises(ValueError, match="Unknown task name"):
        resolve_tasks(
            benchmark="MTEB(Multilingual, v2)",
            task_types=list(DEFAULT_TASK_TYPES),
            task_names=partition,
            allow_missing_task_names=False,
        )

    resolved = resolve_tasks(
        benchmark="MTEB(Multilingual, v2)",
        task_types=types,
        task_names=partition,
        allow_missing_task_names=False,
    )
    assert len(resolved) == len(partition)
    assert {t.metadata.name for t in resolved} == set(partition)


def test_parallel_worker_passes_task_types_to_run_evaluation(tmp_path: Path):
    from mteb_eval.evaluate_parallel import _worker
    from mteb_eval.runner import EvalRunResult

    captured: dict = {}

    def fake_run(args, *, task_names=None, task_types=None, output_dir=None):
        captured["args_task_types"] = list(args.task_types)
        captured["task_types"] = list(task_types) if task_types is not None else None
        captured["task_names"] = list(task_names) if task_names is not None else None
        (Path(output_dir) / "summary.json").write_text("{}", encoding="utf-8")
        return EvalRunResult(
            model_result=ModelResult(
                model_name="dummy",
                model_revision=None,
                task_results=[],
            ),
            timings={},
            failures={},
            task_prompts={},
            exit_code=0,
            model_name="dummy",
        )

    shard = tmp_path / "gpu0"
    shard.mkdir()
    preset_types = [
        "STS",
        "Classification",
        "Clustering",
        "Reranking",
        "Retrieval",
    ]
    args_dict = {
        "cache_dir": None,
        "default_cache": True,
        "model": "dummy",
        "model_path": None,
        "hub_id": None,
        "model_type": "auto",
        "offline": False,
        "benchmark": "MTEB(Multilingual, v2)",
        "task_types": ["STS", "Retrieval"],  # CLI default left in namespace
        "tasks": None,
        "tasks_preset": None,
        "languages": None,
        "languages_preset": "ml16",
        "exclusive_language_filter": False,
        "output_dir": str(tmp_path / "out"),
        "device": None,
        "dtype": "bfloat16",
        "attn_implementation": "sdpa",
        "batch_size": 8,
        "query_batch_size": None,
        "corpus_batch_size": None,
        "overwrite": "only-missing",
        "continue_on_error": True,
        "max_seq_len": 512,
        "query_prefix": None,
        "document_prefix": None,
        "verbose": False,
    }

    with patch("mteb_eval.evaluate_parallel.run_evaluation", side_effect=fake_run):
        code = _worker("0", ["STS12", "MassiveIntentClassification"], args_dict, str(shard), preset_types)

    assert code == 0
    assert captured["task_types"] == preset_types
    assert captured["args_task_types"] == preset_types
    assert captured["task_names"] == ["STS12", "MassiveIntentClassification"]


def test_write_results_workbook_sheets(tmp_path):
    from mteb_eval.summary import write_results_workbook
    from openpyxl import load_workbook

    results = ModelResult(
        model_name="test-model",
        model_revision=None,
        task_results=[
            _sample_task_result("STS12"),
            _sample_task_result("FinancialPhrasebankClassification"),
        ],
    )
    summary_rows = build_summary_rows(
        results,
        {"STS12": 1.0, "FinancialPhrasebankClassification": 2.0},
    )
    path = tmp_path / "results.xlsx"
    write_results_workbook(
        path,
        summary_rows,
        results,
        task_types_by_name={
            "STS12": "STS",
            "FinancialPhrasebankClassification": "Classification",
        },
    )
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "STS" in wb.sheetnames
    assert "Classification" in wb.sheetnames
    assert "ByLanguage" in wb.sheetnames
    assert "Detail" in wb.sheetnames
    assert "Reranking" not in wb.sheetnames  # empty type omitted
